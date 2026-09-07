import io
import pytest
import pandas as pd
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database.models import Base, Material, HarmonizedMaterial, MaterialMapping, MatchResult
from app.utils.validators import validate_dataframe, EXPECTED_COLUMNS
from app.services.ingestion.file_loader import ingest_material_file
from app.services.matching.candidate_search import run_candidate_search_and_matching
from app.services.harmonization.review_service import approve_match
from app.services.export.export_service import generate_harmonized_materials_excel, build_harmonized_materials_dataframe

@pytest.fixture
def test_db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()

def test_missing_column_validation():
    # DataFrame missing required columns like 'Warehouse', 'Supplier', 'Dimensions'
    incomplete_data = {
        "CPSE": ["IOCL"],
        "Legacy Material Code": ["IO-VAL-111"],
        "Raw Material Description": ["GATE VALVE 50MM"]
    }
    df = pd.DataFrame(incomplete_data)
    cleaned_df, report = validate_dataframe(df, check_all_expected=True)

    assert cleaned_df.empty
    assert len(report["errors"]) > 0
    assert len(report["missing_columns"]) > 0
    assert "Warehouse" in report["missing_columns"]
    assert "Supplier" in report["missing_columns"]
    assert "Dimensions" in report["missing_columns"]

def test_full_custom_upload_harmonization_and_export(test_db):
    # 1. Create a custom Excel dataset with all 16 expected columns
    test_records = [
        {
            "CPSE": "IOCL",
            "Legacy Material Code": "TEST-IO-99901",
            "Raw Material Description": "HEX BOLT STAINLESS STEEL M16X50MM ISO 4017",
            "Category": "Fasteners",
            "Subcategory": "Bolts",
            "Material": "Stainless Steel",
            "Specification": "ISO 4017",
            "Dimensions": '{"diameter_mm": 16.0, "length_mm": 50.0, "thread": "M16"}',
            "UOM": "NOS",
            "Unit Price": 120.0,
            "Available Quantity": 500.0,
            "Reserved Quantity": 50.0,
            "On Order Quantity": 100.0,
            "Reorder Level": 100.0,
            "Warehouse": "IOCL Panipat Central Warehouse",
            "Supplier": "Unbrako Fasteners Ltd"
        },
        {
            "CPSE": "BPCL",
            "Legacy Material Code": "TEST-BP-99902",
            "Raw Material Description": "SS HEX HEAD BOLT M16 X 50 MM ISO 4017",
            "Category": "Fasteners",
            "Subcategory": "Bolts",
            "Material": "Stainless Steel",
            "Specification": "ISO 4017",
            "Dimensions": '{"diameter_mm": 16.0, "length_mm": 50.0, "thread": "M16"}',
            "UOM": "NOS",
            "Unit Price": 122.0,
            "Available Quantity": 350.0,
            "Reserved Quantity": 20.0,
            "On Order Quantity": 50.0,
            "Reorder Level": 80.0,
            "Warehouse": "BPCL Mumbai Central Store",
            "Supplier": "TVS Fasteners"
        }
    ]

    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        pd.DataFrame(test_records).to_excel(writer, index=False, sheet_name="Material_Master")
    excel_bytes = bio.getvalue()

    # 2. Ingest custom Excel workbook
    ok, report = ingest_material_file(
        db=test_db,
        file_bytes_or_path=excel_bytes,
        file_name="custom_materials.xlsx",
        actor_name="Test Inspector",
        check_all_expected=True
    )

    assert ok is True
    assert report["valid_records"] == 2
    assert report["new_materials"] == 2

    m1 = test_db.query(Material).filter(Material.legacy_code == "TEST-IO-99901").first()
    m2 = test_db.query(Material).filter(Material.legacy_code == "TEST-BP-99902").first()
    assert m1 is not None
    assert m2 is not None
    assert m1.status == "raw"
    assert len(m1.inventories) == 1
    assert m1.inventories[0].available_qty == 500.0
    assert len(m2.inventories) == 1
    assert m2.inventories[0].available_qty == 350.0

    # 3. Confirm newly loaded materials immediately flow through AI Harmonization pipeline
    matches = run_candidate_search_and_matching(test_db)
    assert len(matches) > 0

    top_match = matches[0]
    assert top_match.final_confidence >= 80.0
    assert top_match.status in ("pending", "approved")

    # 4. Approve candidate match
    hm = approve_match(test_db, top_match.id, officer_name="Test Approver")
    assert hm is not None
    assert hm.status == "approved"
    assert hm.common_code.startswith("NMC-")
    assert len(hm.mappings) == 2

    # 5. Export harmonized materials to Excel
    export_bytes, filename, record_count = generate_harmonized_materials_excel(test_db)
    assert record_count >= 1
    assert filename.startswith("harmonized_materials_export_")
    assert filename.endswith(".xlsx")

    # 6. Validate exported Excel content with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(export_bytes))
    sheet = wb["Harmonized Materials"]
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]

    assert "Common Material Code" in headers
    assert "Standard Description" in headers
    assert "Category" in headers
    assert "Subcategory" in headers
    assert "Material" in headers
    assert "Specification" in headers
    assert "Dimensions" in headers
    assert "UOM" in headers
    assert "IOCL Legacy Code" in headers
    assert "BPCL Legacy Code" in headers
    assert "Total Available Inventory" in headers
    assert "Status" in headers

    # Verify first data row
    cmc_col = headers.index("Common Material Code") + 1
    iocl_col = headers.index("IOCL Legacy Code") + 1
    bpcl_col = headers.index("BPCL Legacy Code") + 1
    avail_col = headers.index("Total Available Inventory") + 1
    status_col = headers.index("Status") + 1

    row_cmc = sheet.cell(2, cmc_col).value
    row_iocl = sheet.cell(2, iocl_col).value
    row_bpcl = sheet.cell(2, bpcl_col).value
    row_avail = sheet.cell(2, avail_col).value
    row_status = sheet.cell(2, status_col).value

    assert row_cmc == hm.common_code
    assert row_iocl == "TEST-IO-99901"
    assert row_bpcl == "TEST-BP-99902"
    assert row_avail == 850.0  # 500 + 350
    assert row_status in ["Approved", "Modified"]
