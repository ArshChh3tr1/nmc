import io
from datetime import datetime
from typing import Tuple, List, Dict, Any
import pandas as pd
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database.models import HarmonizedMaterial, Material, MaterialMapping, Category

def build_harmonized_materials_dataframe(db: Session) -> pd.DataFrame:
    """
    Builds a pandas DataFrame containing all approved/modified harmonized materials
    with structured attributes, individual CPSE legacy code columns, and aggregated inventory.
    """
    # Fetch approved and modified harmonized materials (exclude any pending/rejected)
    harmonized_items = db.query(HarmonizedMaterial).filter(
        HarmonizedMaterial.status.in_(["approved", "modified", "active"])
    ).all()

    # Discover all distinct CPSEs in the database
    db_cpses = [c[0] for c in db.query(Material.cpse).distinct().all() if c[0]]
    canonical_cpses = ["IOCL", "BPCL", "HPCL", "ONGC", "GAIL"]
    all_cpses = sorted(list(set(db_cpses + canonical_cpses)))

    rows: List[Dict[str, Any]] = []

    for hm in harmonized_items:
        # Resolve category & subcategory hierarchy
        cat_name = ""
        subcat_name = ""
        if hm.category:
            if hm.category.parent:
                cat_name = hm.category.parent.name
                subcat_name = hm.category.name
            else:
                cat_name = hm.category.name
                subcat_name = ""

        # Map CPSE legacy codes
        cpse_code_map: Dict[str, List[str]] = {}
        total_available = 0.0

        for mapping in hm.mappings:
            cpse_code_map.setdefault(mapping.cpse, []).append(mapping.legacy_code)
            mat = mapping.material
            if mat and mat.inventories:
                for inv in mat.inventories:
                    total_available += float(inv.available_qty or 0.0)

        # Determine display status
        if str(hm.status).lower() == "modified":
            status_display = "Modified"
        else:
            status_display = "Approved"

        row: Dict[str, Any] = {
            "Common Material Code": hm.common_code,
            "Standard Description": hm.standard_description,
            "Category": cat_name,
            "Subcategory": subcat_name,
            "Material": hm.material_type or "",
            "Specification": hm.specification or "",
            "Dimensions": hm.dimensions_json or "",
            "UOM": hm.uom or "NOS",
        }

        # Individual column per CPSE
        for cpse in all_cpses:
            col_name = f"{cpse} Legacy Code"
            codes = cpse_code_map.get(cpse, [])
            row[col_name] = ", ".join(codes) if codes else ""

        row["Total Available Inventory"] = total_available
        row["Status"] = status_display

        rows.append(row)

    if not rows:
        # Return empty DataFrame with predefined expected columns
        cols = [
            "Common Material Code", "Standard Description", "Category", "Subcategory",
            "Material", "Specification", "Dimensions", "UOM"
        ] + [f"{c} Legacy Code" for c in all_cpses] + ["Total Available Inventory", "Status"]
        return pd.DataFrame(columns=cols)

    return pd.DataFrame(rows)

def generate_harmonized_materials_excel(db: Session) -> Tuple[bytes, str, int]:
    """
    Generates a professionally formatted Excel spreadsheet of all approved/modified
    harmonized materials.
    Returns:
        (excel_bytes, filename, record_count)
    """
    df = build_harmonized_materials_dataframe(db)
    record_count = len(df)
    timestamp_str = datetime.now().strftime("%Y%m%d")
    filename = f"harmonized_materials_export_{timestamp_str}.xlsx"

    bio = io.BytesIO()

    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Harmonized Materials", index=False)
        worksheet = writer.sheets["Harmonized Materials"]

        # Professional enterprise styling
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")

        # Format header row
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        worksheet.row_dimensions[1].height = 28

        # Format data rows and apply zebra shading
        for row_idx in range(2, len(df) + 2):
            worksheet.row_dimensions[row_idx].height = 20
            is_even = (row_idx % 2 == 0)
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

                col_name = df.columns[col_idx - 1]
                if col_name == "Total Available Inventory":
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0.00"
                elif col_name in ["Common Material Code", "UOM", "Status"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths with safety margins
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    bio.seek(0)
    return bio.getvalue(), filename, record_count
